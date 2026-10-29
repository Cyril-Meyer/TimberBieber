"""
Programme principal : produit un seul fichier  results/resultats_<matière>.xlsx

    python run.py

Feuilles :
  1. Production par semaine : quantité à couper par article et par semaine
  2. Ordre de coupe         : panneaux dans l'ordre de coupe, avec les OF et leurs dates
  3. Comparaison Schmidt    : les 5 graphiques Schmidt originaux (gauche) et les nôtres
                              au même format et aux mêmes semaines (droite) + tableau
                              + tableau des valeurs

Méthode retenue : lots hebdomadaires (OF = besoin de la semaine suivante + stock de
sécurité, arrondi à QteParOF), imbrication de toute la semaine en une fois,
remplissage des chutes avec des pièces nécessaires dans les 3 semaines au plus.
"""
import os
import tempfile

import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment

from .data import load_material, weekly_demand, _p, F_ART
from .nesting import BestNester, Panel
from .planning import params, simulate
from .sequencing import cutting_order, LIMIT
from . import comparaison

MATIERE = 215371
SS_WEEKS = 0.5            # stock de sécurité au montage (semaines de demande)
SMAX_WEEKS = 3.0          # stock maximum au montage (semaines de demande)
LOT_RULE = "min"          # "min" : OF >= QteParOF ; "multiple" : OF = k x QteParOF
MAX_DELAI_SEMAINES = 3    # délai max. entre lancement de l'OF et utilisation des pièces
SEUIL_VOLUME = 0.5        # comparaison : semaine ignorée si moins de 50 % du volume médian
                          # de la matière (historique OF incomplet, ex. début de l'extraction)
RESTES_AGE_MAX = 4        # un reste (>= ResteMin) est gardé 4 semaines au plus, puis mis au rebut
OUT = "results"


def _label(w):
    iso = w.isocalendar()
    return f"{iso[0]}/S{iso[1]:02d}"


def _style(ws, widths, freeze="B2"):
    head = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = freeze


def run():
    os.makedirs(OUT, exist_ok=True)
    mat, art = load_material(MATIERE)
    D = weekly_demand(art)
    panel = Panel(L=int(mat.X), W=int(mat.Y),
                  reste_min_l=int(mat.ResteMin_X), reste_min_w=int(mat.ResteMin_Y))
    nester = BestNester({a: (r.X, r.Y, r.SensFil) for a, r in art.iterrows()}, panel)
    p = params(art, D, ss_weeks=SS_WEEKS, smax_weeks=SMAX_WEEKS)

    # ---------- planification hebdomadaire
    kpi, X = simulate(p, D, nester, "policy", fill=True, lot_rule=LOT_RULE,
                      max_weeks=MAX_DELAI_SEMAINES, restes_age_max=RESTES_AGE_MAX)
    tdc = 1 - kpi.piece_area.sum() / kpi.consumed_area.sum()
    print(f"matière {MATIERE} : {len(X)} semaines, TdC = {tdc:.1%}, "
          f"{int(kpi.panels.sum())} panneaux, stock min. {kpi.min_stock_after.min():.0f}")

    # ---------- feuille 1 : production par article et par semaine
    prod = X.T.astype(int)
    prod.columns = [_label(w) for w in X.index]
    prod.insert(0, "Désignation", art["Désignation article"])
    prod.insert(1, "QteParOF", p.q)
    prod["Total"] = prod.iloc[:, 2:].sum(axis=1)
    prod = prod.sort_values("Total", ascending=False)
    prod.index.name = "Article"
    tot = pd.DataFrame([["TOTAL pièces", ""] + list(prod.iloc[:, 2:].sum())],
                       columns=prod.columns, index=[""])
    tdc_row = pd.DataFrame([["Taux de chute", ""] + [round(v * 100, 1) if n else None
                            for v, n in zip(kpi.tdc, kpi.panels)] + [round(tdc * 100, 1)]],
                           columns=prod.columns, index=[""])
    pan_row = pd.DataFrame([["Panneaux", ""] + list(kpi.panels.astype(int)) + [int(kpi.panels.sum())]],
                           columns=prod.columns, index=[""])
    res_row = pd.DataFrame([["Restes réutilisés", ""] + list(kpi.restes_utilises.astype(int))
                            + [int(kpi.restes_utilises.sum())]], columns=prod.columns, index=[""])
    sheet1 = pd.concat([prod, tot, pan_row, res_row, tdc_row])
    sheet1.index.name = "Article"

    # ---------- feuille 2 : ordre de coupe (avec les dates des OF)
    desig = art["Désignation article"].to_dict()
    ordre, ofs, peaks = [], [], []
    for w, row in kpi.iterrows():
        x = row["_x"][row["_x"] > 0].to_dict()
        r_ordre, r_of, peak = cutting_order(w, row["_panels"], x, p.q, desig)
        ordre += r_ordre; ofs += r_of; peaks.append(peak)
    ofs = pd.DataFrame(ofs)
    assert (ofs["Délai (jours)"] <= 7 * MAX_DELAI_SEMAINES).all()
    sheet2 = pd.DataFrame(ordre).merge(
        ofs[["OF", "Quantité", "Date lancement", "Date livraison"]], on="OF", how="left")
    sheet2 = sheet2.rename(columns={"Quantité": "Taille OF"})
    sheet2["Semaine"] = sheet2["Semaine"].str.replace("-", "/")
    print(f"{len(ofs)} OF, pic du stock tampon {max(peaks):.0f} (limite {LIMIT})")

    # ---------- feuille 3 : comparaison avec Schmidt (graphiques côte à côte)
    tmp = tempfile.mkdtemp()
    mat_all = pd.read_excel(_p(F_ART), "Matière").dropna(subset=["Matiere"])
    graphs = [("global", "Taux de Chute – toutes matières (épaisseur 16)",
               comparaison.weekly_tdc_global(mat_all.Matiere.astype(int), RESTES_AGE_MAX))]
    for m in (215371, 221612, 220965, 219106):
        graphs.append((m, f"TdC par Coloris – {m}", comparaison.weekly_tdc(m, RESTES_AGE_MAX)))
    pairs, tables = [], []
    for nom, titre, ours in graphs:
        sch = comparaison.SCHMIDT[nom]
        weeks = list(sch)                          # uniquement les semaines Schmidt
        # semaines avec trop peu d'OF dans l'historique -> ignorées
        seuil = SEUIL_VOLUME * ours.pieces.median()
        garde = ours.index[ours.pieces >= seuil]
        ys = [ours.B0w[w] * 100 if w in garde else np.nan for w in weeks]
        y0 = [ours.B0[w] * 100 if w in garde else np.nan for w in weeks]
        pcs = [int(ours.pieces[w]) if w in ours.index else 0 for w in weeks]
        statut = ["pas d'OF dans le fichier" if w not in ours.index
                  else ("ignorée : historique insuffisant" if w not in garde else "")
                  for w in weeks]
        img_s = comparaison.schmidt_image(nom, os.path.join(tmp, f"schmidt_{nom}.png"))
        img_n = os.path.join(tmp, f"nous_{nom}.png")
        # semaines sans donnée retirées de l'axe -> courbe continue (comme Schmidt)
        k = [i for i, v in enumerate(ys) if not np.isnan(v)]
        comparaison.chart_nous(nom, [weeks[i] for i in k], [ys[i] for i in k],
                               [y0[i] for i in k], img_n, titre)
        pairs.append((titre, img_s, img_n))
        ok = ~np.isnan(ys)
        tables.append(pd.DataFrame({
            "Graphique": titre, "Semaine": [f"2026/S{w}" for w in weeks],
            "TdC Schmidt (%)": [sch[w] for w in weeks],
            "Pièces dans l'historique": pcs, "Semaine retenue ?": statut,
            "Notre recalcul, OF coupés par jour (%)": np.round(y0, 2),
            "OF regroupés par semaine (%)": np.round(ys, 2)}))
        print(f"{titre:45s} Schmidt {np.mean(np.array([sch[w] for w in weeks])[ok]):5.2f} %"
              f" | recalcul par jour {np.nanmean(y0):5.2f} % | par semaine {np.nanmean(ys):5.2f} %")
    sheet3 = pd.concat(tables, ignore_index=True)

    # ---------- écriture
    path = os.path.join(OUT, f"resultats_{MATIERE}.xlsx")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        sheet1.to_excel(xw, sheet_name="Production par semaine")
        sheet2.to_excel(xw, sheet_name="Ordre de coupe", index=False)

        ws = xw.sheets["Production par semaine"]
        _style(ws, {"A": 10, "B": 32, "C": 10}, freeze="D2")
        for row in ws.iter_rows(min_row=ws.max_row - 3, max_row=ws.max_row):
            for c in row:
                c.font = Font(bold=True)

        ws = xw.sheets["Ordre de coupe"]
        _style(ws, {"A": 10, "B": 11, "C": 9, "D": 32, "E": 20, "F": 12, "G": 10,
                    "H": 12, "I": 9, "J": 14, "K": 14}, freeze="A2")
        ws.auto_filter.ref = ws.dimensions

        # feuille 3 : graphiques Schmidt (gauche) et les nôtres (droite), puis le tableau
        ws = xw.book.create_sheet("Comparaison Schmidt")
        ws["A1"] = "Schmidt (original)"; ws["J1"] = "Notre calcul (mêmes OF, mêmes semaines)"
        ws["A1"].font = ws["J1"].font = Font(bold=True, size=13)
        W, H, STEP = 960, 304, 18                  # taille des images (px), lignes par bloc
        for i, (titre, img_s, img_n) in enumerate(pairs):
            r = 3 + i * STEP
            ws[f"A{r - 1}"] = titre
            ws[f"A{r - 1}"].font = Font(bold=True)
            for img, col in ((img_s, "A"), (img_n, "J")):   # A–E = 774 px, J commence à ~1030 px
                im = XLImage(img); im.width, im.height = W, H
                ws.add_image(im, f"{col}{r}")
        r0 = 3 + len(pairs) * STEP + 1
        ws.cell(r0 - 1, 1, "Valeurs. Semaines ignorées : pas d'OF dans le fichier fourni, ou moins "
                           "de 50 % du volume médian de la matière (historique incomplet). "
                           "Valeurs Schmidt lues sur leurs graphiques ; moyennes sur les semaines retenues.").font = Font(italic=True)
        for j, col in enumerate(sheet3.columns, start=1):
            c = ws.cell(r0, j, col); c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for i, row in enumerate(sheet3.itertuples(index=False), start=1):
            for j, v in enumerate(row, start=1):
                ws.cell(r0 + i, j, None if isinstance(v, float) and np.isnan(v) else v)
        for col, w in {"A": 44, "B": 11, "C": 14, "D": 14, "E": 18}.items():
            ws.column_dimensions[col].width = w
    print("écrit :", path)
    return path


if __name__ == "__main__":
    run()
